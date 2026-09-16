import tkinter as tk

from tkinter import PhotoImage

from tkinter import ttk, messagebox

from tkcalendar import DateEntry

from pymongo import MongoClient

from bson.objectid import ObjectId

import pandas as pd

import datetime

import bcrypt

import sys

import os

from PIL import Image, ImageTk

import webbrowser


from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from reportlab.platypus.paragraph import ParagraphStyle

from reportlab.platypus.flowables import KeepTogether

from reportlab.lib import colors

from reportlab.lib.styles import getSampleStyleSheet

from reportlab.lib.pagesizes import A4

from reportlab.pdfgen import canvas

from reportlab.platypus import Image as RLImage

from openpyxl import Workbook

from openpyxl.styles import Alignment, Font



# ====== DATABASE CONNECTION ======


client = MongoClient("mongodb://localhost:27017/")


database = client["ControleDeEstoque"]


products_collection = database["produtos"]

movements_collection = database["movimentacao"]

users_collection = database["usuarios"]


MOVEMENT_TYPE_LABELS = {"ENTRADA": "INBOUND", "SAÍDA": "OUTBOUND"}
MOVEMENT_TYPE_VALUES = {label: value for value, label in MOVEMENT_TYPE_LABELS.items()}

def display_movement_type(value):
    return MOVEMENT_TYPE_LABELS.get(value, value)

def persisted_movement_type(value):
    return MOVEMENT_TYPE_VALUES.get(value.upper(), value)




# ====== ADMINISTRATOR USER CREATION ======

def create_admin_user():
    if users_collection.count_documents({"username":"a"}) == 0:

        password = "a"

        hashed = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt())


        users_collection.insert_one({
            "username": "a",
            "password": hashed,
            "plain_password": password
        })


# ====== LOGIN VERIFICATION ======

def verify_login():

    username = username_entry.get()

    password = password_entry.get()



    if not username or not password:

        messagebox.showerror("Error", "Please fill in all fields")

        return


    user = users_collection.find_one({"username":username})


    if user and bcrypt.checkpw(password.encode('utf-8'), user["password"]):

        login_window.destroy()



        open_main_window()


    else:

        messagebox.showerror("Error", "Incorrect username or password")






# =============== USER MANAGEMENT SCREEN =====================


def open_user_management():

    user_window = tk.Toplevel()

    user_window.title("Manage Users")

    user_window.geometry("800x600")

    user_window.resizable(False, False)

    user_window.update_idletasks()

    width  = user_window.winfo_width()

    height = user_window.winfo_height()


    pos_x = (user_window.winfo_screenwidth() // 2 ) - (width // 2 )

    pos_y = (user_window.winfo_screenheight() // 2 ) - (height // 2 )


    user_window.geometry(f"{width}x{height}+{pos_x}+{pos_y}")



    # ====== LOAD USER DATA INTO THE TREEVIEW ======

    def load_users():

        for item in users_tree.get_children():

            users_tree.delete(item)


        records = users_collection.find()


        for doc in records:

            users_tree.insert(
                "",
                "end",
                values=(
                    str(doc["_id"]),
                    doc["username"],
                    doc.get("plain_password", "")

                )

            )


    # ====== USER WINDOW TOP FRAME ======
    top_frame = tk.Frame(user_window, pady=10)

    top_frame.pack()



    tk.Label(
        top_frame,
        text="Manage Users",
        font=("Helvetica", 16, "bold")).pack()





    form_frame = tk.Frame(user_window, padx=20, pady=10)

    form_frame.pack()

    tk.Label(
        form_frame,
        text="Username",
        font=("Helvetica", 12)
    ).grid(row=0, column=0, sticky="e")


    new_username_entry = tk.Entry(
        form_frame,
        font=("Helvetica", 12))

    new_username_entry.grid(row=0, column=1, padx=10, pady=5 )




    tk.Label(
        form_frame,
        text="Password",
        font=("Helvetica", 12)
    ).grid(row=1, column=0, sticky="e")


    new_password_entry = tk.Entry(
        form_frame,
        font=("Helvetica", 12))

    new_password_entry.grid(row=1, column=1, padx=10, pady=5 )


    # ====== BUTTON FRAME ======

    button_frame = tk.Frame(user_window, pady=10)

    button_frame.pack()


    # ====== ADD USER BUTTON ======

    def add_user():

        username = new_username_entry.get()

        password = new_password_entry.get()


        if not username or not password:

            messagebox.showerror("Error", "Please fill in all fields.")

            return


        if users_collection.find_one({"username": username}):

            messagebox.showerror("Error", "This username already exists in the database.")

            return


        hashed = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt())


        users_collection.insert_one({
            "username": username,
            "password": hashed,
            "plain_password": password
        })

        messagebox.showinfo("Success", "User registered successfully.")


        new_username_entry.delete(0, tk.END)

        new_password_entry.delete(0, tk.END)


        load_users()


        user_window.destroy()




    add_user_button = tk.Button(
        button_frame,
        text="Add User",
        command=add_user,
        bg="#CDC9C9",
        fg="black",
        font=("Helvetica", 12),
        width=15
    )

    add_user_button.pack(side="left", padx=5)



    add_user_button.bind(
    "<Enter>",
    lambda e: add_user_button.config(fg="white", bg="#8B8989"))



    add_user_button.bind(
    "<Leave>",
    lambda e: add_user_button.config(fg="black", bg="#CDC9C9"))



    # ====== EDIT USER BUTTON ======

    def edit_user():

        selected = users_tree.selection()

        if not selected:

            messagebox.showerror("Error", "Select a user")

            return


        item = users_tree.item(selected)

        doc_id = item["values"][0]


        username = new_username_entry.get()

        new_password = new_password_entry.get()



        if not username or not new_password:

            messagebox.showerror("Error", "Please fill in all fields.")

            return

        hashed = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt())


        users_collection.update_one(
            {"_id": ObjectId(doc_id)},
            {"$set":{
                "username": username,
                "password" : hashed,
                "plain_password": new_password
            }}
        )


        messagebox.showinfo("Success", f"User '{username}' updated successfully.")


        new_username_entry.delete(0, tk.END)

        new_password_entry.delete(0, tk.END)


        load_users()



    edit_user_button = tk.Button(
        button_frame,
        text="Edit User",
        command=edit_user,
        bg="#CDC9C9",
        fg="black",
        font=("Helvetica", 12),
        width=15
    )

    edit_user_button.pack(side="left", padx=5)



    edit_user_button.bind(
    "<Enter>",
    lambda e: edit_user_button.config(fg="white", bg="#8B8989"))



    edit_user_button.bind(
    "<Leave>",
    lambda e: edit_user_button.config(fg="black", bg="#CDC9C9"))



    # ====== DELETE USER BUTTON ======

    def delete_user():

        selected = users_tree.selection()

        if not selected:

            messagebox.showerror("Error", "Select a user to delete.")


            return


        item = users_tree.item(selected)

        username = item["values"][1]

        doc_id = item["values"][0]


        if messagebox.askyesno("Confirmation", f"Are you sure you want to delete user '{username}'?"):

            users_collection.delete_one({"_id": ObjectId(doc_id)})


            messagebox.showinfo("Success", f"User '{username}' deleted successfully.")

            load_users()


    delete_user_button = tk.Button(
        button_frame,
        text="Delete User",
        command=delete_user,
        bg="#CDC9C9",
        fg="black",
        font=("Helvetica", 12),
        width=15
    )

    delete_user_button.pack(side="left", padx=5)



    delete_user_button.bind(
    "<Enter>",
    lambda e: delete_user_button.config(fg="white", bg="#8B8989"))



    delete_user_button.bind(
    "<Leave>",
    lambda e: delete_user_button.config(fg="black", bg="#CDC9C9"))


    # ====== USER TABLE TREEVIEW ======

    def populate_user_fields(event):

        selected = users_tree.selection()

        if selected:

            item = users_tree.item(selected)

            values = item['values']


            new_username_entry.delete(0, tk.END)

            new_username_entry.insert(0, values[1])


            new_password_entry.delete(0, tk.END)

            new_password_entry.insert(0, values[2])



    list_frame = tk.Frame(user_window, pady=10)

    list_frame.pack(fill="both", expand=True)



    users_tree = ttk.Treeview(
        list_frame,
        columns=("ID", "Username", "Password"),
        show="headings"
    )


    users_tree.heading("ID", text="ID")

    users_tree.column("ID", width=100, anchor="center")




    users_tree.heading("Username", text="Username")

    users_tree.column("Username", width=200, anchor="center")




    users_tree.heading("Password", text="Password")

    users_tree.column("Password", width=200, anchor="center")



    users_tree.pack(fill="both", expand=True)


    users_tree.bind("<<TreeviewSelect>>", populate_user_fields)

    load_users()












# REMOVE PRODUCT QUANTITY FROM STOCK

def remove_product_stock():

    selected = products_tree.selection()

    if not selected:

        messagebox.showerror("Error", "Select a product to remove from stock")

        return



    item = products_tree.item(selected)

    doc_id = item["values"][5]

    product_code = item["values"][0]

    product_name = item["values"][1]

    current_quantity = item["values"][3]

    removal_quantity = quantity_entry.get()


    if not removal_quantity:

        messagebox.showerror("Error", "Enter the quantity to remove.")

        return

    try:

        removal_quantity = int(removal_quantity)

    except ValueError:

        messagebox.showerror("Error", "The quantity to remove must be a number.")

        return

    if removal_quantity > current_quantity:

        messagebox.showerror("Error", "The requested quantity exceeds the available stock.")

        return

    new_quantity = current_quantity - removal_quantity


    products_collection.update_one(
        {"_id": ObjectId(doc_id)},
        {"$set": {"quantidade": new_quantity}}
        )


    record_movement(product_code, product_name, "SAÍDA", removal_quantity)

    messagebox.showinfo("Success", f"Removed {removal_quantity} units successfully.")


    load_data()


    clear_fields()

















# ADD A NEW PRODUCT TO THE SYSTEM

def add_product():

    code = code_entry.get().strip().upper()

    name = name_entry.get().strip().upper()

    price = price_entry.get().strip().upper()

    quantity = quantity_entry.get().strip().upper()

    location = location_entry.get().strip().upper()



    if not code or not  name or not price or not quantity or not location:

        messagebox.showerror("Error", "All fields are required")

        return


    try:

        price = float(price)

        quantity = int(quantity)


    except ValueError:

        messagebox.showerror("Error", "Price and quantity must be numeric values.")

        return


    products_collection.insert_one({
        "codigo": code,
        "nome" : name,
        "preco": price,
        "quantidade":quantity,
        "localizacao": location

    })



    record_movement(code, name, "ENTRADA", quantity)

    messagebox.showinfo("Success", "Product registered successfully.")



    load_data()

    clear_fields()




def record_movement(code, product, movement_type, quantity):

    current_date = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    found_product = products_collection.find_one({"codigo": code})

    movements_collection.insert_one({

        "codigo": code,
        "produto": product,
        "tipo":movement_type,
        "quantidade":quantity,
        "data": current_date,
        "localizacao": found_product.get("localizacao", "")

    })


def clear_fields():

    code_entry.delete(0, tk.END)

    name_entry.delete(0, tk.END)

    price_entry.delete(0, tk.END)

    quantity_entry.delete(0, tk.END)

    location_entry.delete(0, tk.END)



# LOAD PRODUCT DATA INTO THE TREEVIEW

def load_data():

    for item in products_tree.get_children():

        products_tree.delete(item)



    records = products_collection.find({"quantidade": {"$gt": 0}})

    for doc in records:

        quantity = doc.get("quantidade", 0)


        item_id = products_tree.insert(
            "",
            "end",
            values=(
                doc.get("codigo", ""),
                doc.get("nome", ""),
                doc.get("preco", ""),
                quantity,
                doc.get("localizacao", ""),
                str(doc.get("_id", ""))
            )
        )



# FILTER PRODUCTS BY PRODUCT OR LOCATION
def filter_product_or_location():

    term = filter_entry.get().strip()

    if not term:

        load_data()

        return

    for item in products_tree.get_children():

        products_tree.delete(item)


    regex = {"$regex" : term, "$options" : "i"}


    try:

        price_value = float(term)

        quantity_value = int(term)


    except ValueError:

        price_value = None

        quantity_value = None



    query = {"$or": [
        {"nome" : regex},
        {"localizacao" : regex}
    ]}


    if price_value is not None:

        query["$or"].append({"preco": price_value})


    if quantity_value is not None:

        query["$or"].append({"quantidade": quantity_value})



    records = products_collection.find(query)


    for doc in records:

        quantity = doc["quantidade"]

        item_id = products_tree.insert(
            "",
            "end",
            values=(
                doc.get("codigo", ""),
                doc.get("nome", ""),
                doc.get("preco", ""),
                quantity,
                doc.get("localizacao", ""),
                str(doc.get("_id", ""))

            )

        )













# POPULATE FIELDS WITH THE SELECTED PRODUCT DATA


def populate_product_fields(event):

    selected = products_tree.selection()

    if selected:

        item = products_tree.item(selected)

        values = item['values']


        code_entry.delete(0, tk.END)
        code_entry.insert(0, values[0])


        name_entry.delete(0, tk.END)
        name_entry.insert(0, values[1])


        price_entry.delete(0, tk.END)
        price_entry.insert(0, values[2])



        quantity_entry.delete(0, tk.END)
        quantity_entry.insert(0, values[3])



        location_entry.delete(0, tk.END)
        location_entry.insert(0, values[4])


        quantity_entry.focus_set()












# ====== TEMPORARY FUNCTIONS ======



def show_movements():
    print("Waiting for show_movements implementation")






def delete_user():
    print("Waiting for delete_user implementation")



# ====== END TEMPORARY FUNCTIONS ======




# ====== MAIN APPLICATION WINDOW ======

def  open_main_window():

    global main_window

    main_window = tk.Tk()

    main_window.title("Inventory Management System")

    main_window.state('zoomed')

    main_window.resizable(True, True)



    # ------------------------ TOP FRAME ----------------------------


    top_frame = tk.Frame(main_window, bg="#E8E8E8", height=70)

    top_frame.pack(fill="x")


    title_label = tk.Label(top_frame,
                          text="Inventory Management System",
                          bg="#E8E8E8",
                          fg="black",
                          font=("Helvetica", 20, "bold"))


    title_label.pack(pady=10)




    # ------------------------- DATA FRAME -----------------------------


    data_frame = tk.Frame(main_window, padx=20, pady=10)

    data_frame.pack(fill="x")

    # -------------------------------------------------------------------

    tk.Label(data_frame,
             text="Code",
             font=("Helvetica", 12)).grid(row=0, column=0, sticky="e")


    global code_entry


    code_entry = tk.Entry(data_frame, font=("Helvetica", 12))

    code_entry.grid(row=0, column=1, padx=10, pady=5)

    # -------------------------------------------------------------------

    tk.Label(data_frame,
             text="Name",
             font=("Helvetica", 12)).grid(row=0, column=2, sticky="e")

    global name_entry

    name_entry = tk.Entry(data_frame, font=("Helvetica", 12))

    name_entry.grid(row=0, column=3, padx=10, pady=5)


    # -------------------------------------------------------------------

    tk.Label(data_frame,
             text="Price",
             font=("Helvetica", 12)).grid(row=0, column=4, sticky="e")

    global price_entry

    price_entry = tk.Entry(data_frame, font=("Helvetica", 12))

    price_entry.grid(row=0, column=5, padx=10, pady=5)


    # -------------------------------------------------------------------

    tk.Label(data_frame,
             text="Quantity",
             font=("Helvetica", 12)).grid(row=0, column=6, sticky="e")

    global quantity_entry

    quantity_entry = tk.Entry(data_frame, font=("Helvetica", 12))

    quantity_entry.grid(row=0, column=7, padx=10, pady=5)


    # -------------------------------------------------------------------

    tk.Label(data_frame,
             text="Location",
             font=("Helvetica", 12)).grid(row=0, column=8, sticky="e")

    global location_entry

    location_entry = tk.Entry(data_frame, font=("Helvetica", 12))

    location_entry.grid(row=0, column=9, padx=10, pady=5)



    # ========================== MOVEMENT HISTORY WINDOW ==========================
    def show_movements():

        movements_window = tk.Toplevel(main_window)

        movements_window.title("Movement History")

        movements_window.resizable(False, False)

        movements_window.update_idletasks()

        width = 1500

        height = 600

        pos_x = (movements_window.winfo_screenwidth() // 2) - (width // 2)

        pos_y = (movements_window.winfo_screenheight() // 2) - (height // 2)


        movements_window.geometry(f"{width}x{height}+{pos_x}+{pos_y}")



        # ------------------------- MOVEMENT HISTORY FILTER FRAME -----------------------------
        filters_frame = tk.Frame(movements_window, pady=10)

        filters_frame.pack()



        tk.Label(filters_frame, text="Movement History",
                 font=("arial", 22, "bold")).grid(row=0, column=0, columnspan=8, padx=(500, 0), pady=(10, 30))



        tk.Label(filters_frame, text="Code",
                 font=("arial", 12)).grid(row=1, column=0, padx=5)


        code_filter_entry = tk.Entry(filters_frame,
                                         font=("Helvetica", 12))

        code_filter_entry.grid(row=1, column=1, padx=5)






        tk.Label(filters_frame, text="Product",
                 font=("arial", 12)).grid(row=1, column=2, padx=5)


        product_filter_entry = tk.Entry(filters_frame,
                                         font=("Helvetica", 12))

        product_filter_entry.grid(row=1, column=3, padx=5)





        tk.Label(filters_frame, text="Type",
                 font=("arial", 12)).grid(row=1, column=4, padx=5)


        type_filter_entry = tk.Entry(filters_frame,
                                         font=("Helvetica", 12))

        type_filter_entry.grid(row=1, column=5, padx=5)





        tk.Label(filters_frame, text="Start Date",
                 font=("arial", 12)).grid(row=1, column=6, padx=5)


        start_date_entry = DateEntry(filters_frame,
                                         font=("Helvetica", 12),
                                         date_pattern = 'yyyy-MM-dd')

        start_date_entry.grid(row=1, column=7, padx=5)


        # ------------ FIND THE OLDEST MOVEMENT AUTOMATICALLY ------------

        first_movement = movements_collection.find_one(
            {},
            sort=[("data", 1)]
        )


        if first_movement:

            first_date = datetime.datetime.strptime(
                first_movement["data"],
                "%Y-%m-%d %H:%M:%S"
            ).date()

            start_date_entry.set_date(first_date)





        tk.Label(filters_frame,
                 text="End Date:",
                 font=("arial", 12)).grid(row=1, column=8, padx=5)


        end_date_entry = DateEntry(filters_frame,
                                         font=("Helvetica", 12),
                                         date_pattern = 'yyyy-MM-dd')

        end_date_entry.grid(row=1, column=9, padx=5)



        # ------------------ APPLY FILTER BUTTON ------------------

        def load_movements():

            for item in movements_tree.get_children():

                movements_tree.delete(item)


            filter_query = {}


            code = code_filter_entry.get().strip()

            product = product_filter_entry.get().strip()

            movement_type = persisted_movement_type(type_filter_entry.get().strip())

            start_date = start_date_entry.get_date()

            end_date = end_date_entry.get_date()


            if code:

                filter_query["codigo"] = {"$regex": code, "$options": "i"}


            if product:

                filter_query["produto"] = {"$regex": product, "$options": "i"}


            if movement_type:

                filter_query["tipo"] = {"$regex": movement_type, "$options": "i"}



            if start_date or end_date:

                filter_query["data"] = {}

                if start_date:

                    filter_query["data"]["$gte"] = start_date.strftime("%Y-%m-%d")


                if end_date:

                    filter_query["data"]["$lte"] = end_date.strftime("%Y-%m-%d") + "23:59:59"


            records = movements_collection.find(filter_query).sort("data", 1)

            total = 0


            for doc in records:

                movements_tree.insert("", "end",
                                values=(
                                    doc.get("codigo", ""),
                                    doc.get("produto", ""),
                                    display_movement_type(doc.get("tipo", "")),
                                    doc.get("quantidade", ""),
                                    doc.get("data", ""),
                                    doc.get("localizacao", "")))

                total += 1

            total_records_label.config(text=f"Total records: {total}")





        def clear_filters():

            code_filter_entry.delete(0, tk.END)

            product_filter_entry.delete(0, tk.END)

            type_filter_entry.delete(0, tk.END)


            load_movements()





        # EXPORT MOVEMENT HISTORY TO AN EXCEL FILE
        def export_movements():

            filter_query = {}


            code = code_filter_entry.get().strip()

            product = product_filter_entry.get().strip()

            movement_type = persisted_movement_type(type_filter_entry.get().strip())



            start_date = start_date_entry.get_date()

            end_date = end_date_entry.get_date()


            if code:
                filter_query["codigo"] = {"$regex": code, "$options": "i"}


            if product:
                filter_query["produto"] = {"$regex": product, "$options": "i"}


            if movement_type:
                filter_query["tipo"] = {"$regex": movement_type, "$options": "i"}




            if start_date or end_date:

                filter_query["data"] = {}


                if start_date:

                    filter_query["data"]["$gte"] = start_date.strftime("%Y-%m-%d")

                if end_date:

                    filter_query["data"]["$lte"] = end_date.strftime("%Y-%m-%d") + "23:59:59"


            records = list(movements_collection.find(filter_query).sort("data", 1))


            if not records:

                messagebox.showinfo("Export", "There are no records to export")

                return


            wb = Workbook()

            ws = wb.active

            ws.title="Movements"

            headers = ["CODE", "PRODUCT", "TYPE", "QUANTITY", "LOCATION", "DATE"]

            ws.append(headers)

            for col in ws[1]:

                col.font= Font(bold=True)

                col.alignment = Alignment(horizontal="center", vertical="center")


            for doc in records:

                row_data = [
                    str(doc.get("codigo", "")).upper(),
                    str(doc.get("produto", "")).upper(),
                    str(display_movement_type(doc.get("tipo", ""))).upper(),
                    str(doc.get("quantidade", "")).upper(),
                    str(doc.get("localizacao", "")).upper(),
                    str(doc.get("data", "")).upper()
                ]

                ws.append(row_data)


            for row in ws.iter_rows():

                for cell in row:

                    cell.alignment = Alignment(horizontal="center", vertical="center")


            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 30


            filename = "filtered_movements.xlsx"

            wb.save(filename)


            messagebox.showinfo("Export", f"Data exported successfully to: {filename}")




        # FORMAT VALUES AS BRAZILIAN CURRENCY
        def format_brl_currency(value):
            return f"R${value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")



        # GENERATE AN INVENTORY REPORT
        def generate_inventory():

            filter_query = {}

            code = code_filter_entry.get().strip()

            product = product_filter_entry.get().strip()

            movement_type = persisted_movement_type(type_filter_entry.get().strip())



            if code:
                filter_query["codigo"] = {"$regex": code, "$options": "i"}


            if product:
                filter_query["produto"] = {"$regex": product, "$options": "i"}


            if movement_type:
                filter_query["tipo"] = {"$regex": movement_type, "$options": "i"}




            records = list(
                products_collection.find(filter_query).sort("codigo", 1)

            )


            if not records:

                messagebox.showinfo("Inventory", "No products were found for the inventory report.")
                return


            folder = "INVENTORIES"

            os.makedirs(folder, exist_ok=True)

            date_time = datetime.datetime.now().strftime("%d-%m-%Y_%H-%M-%S")


            filename = f"inventory_{date_time}.pdf"


            pdf_path = os.path.join(folder, filename)



            doc = SimpleDocTemplate(
                pdf_path,
                pagesize=A4,
                topMargin=25,
                bottomMargin=25,
                leftMargin=30,
                rightMargin=30
            )


            elements = []

            styles = getSampleStyleSheet()



            title_style = ParagraphStyle(
                "ERPTitle",
                parent=styles["title"],
                alignment=0,
                spaceBefore=0,
                spaceAfter=6,
                leftIndent=40
            )


            subtitle_style = ParagraphStyle(
                "ERPSubtitle",
                parent=styles["Heading4"],
                alignment=0,
                spaceBefore=0,
                spaceAfter=4,
                leftIndent=20
            )


            def absolute_path(relative_path):

                if hasattr(sys, "_MEIPASS"):
                    return os.path.join(sys._MEIPASS, relative_path)


                return os.path.join(os.path.abspath("."), relative_path)



            logo_path = absolute_path("logo.png")

            logo = None

            if os.path.exists(logo_path):

                logo = RLImage(logo_path, width=70, height=70)


            title = Paragraph(
                "<b>INVENTORY REPORT</b>",
                title_style
            )

            header_data = [[logo, title]]

            header_table = Table(
                header_data,
                colWidths=[80, 400]

            )


            header_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))

            elements.append(header_table)
            elements.append(Spacer(1, 15))


            report_name = f"Inventory - {date_time}"

            elements.append(
                Paragraph(f"<b>{report_name}</b>", subtitle_style )

            )

            elements.append(Spacer(1, 10))



            data = [
                ["Code", "Product", "Quantity", "Unit Price", "Total"]
            ]


            grand_total = 0



            for db_product in records:

                code = db_product.get("codigo", "")
                product_name = db_product.get("nome", "")
                quantity = db_product.get("quantidade", 0)
                price = float(db_product.get("preco", 0))


                item_total = price * quantity


                grand_total += item_total




                data.append([
                    code,
                    product_name,
                    quantity,
                    format_brl_currency(price),
                    format_brl_currency(item_total)
                ])


            elements.append(
                Paragraph(
                    f"<b>TOTAL INVENTORY VALUE: {format_brl_currency(grand_total)}</b>",
                    subtitle_style
                )
            )

            elements.append(Spacer(1, 15))


            table = Table(data, colWidths=[70, 180, 80, 90, 90 ])

            table.repeatRows = 1

            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5,  colors.black),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("ALIGN", (2, 1), (-1, -1), "CENTER"),
                ("PADDING", (0, 0), (-1, -1), 6),

            ]))

            elements.append(table)


            doc.build(elements)

            messagebox.showinfo("Success", "Inventory report generated successfully.")


            absolute_pdf_path = os.path.abspath(pdf_path)

            webbrowser.open(f"file://{absolute_pdf_path}")





































        apply_filters_button = tk.Button(
            filters_frame,
            text="Apply Filters",
            command=load_movements,
            bg="#CDC9C9",
            fg="black",
            font=("Helvetica", 12)
        )

        apply_filters_button.grid(row=1, column=10, padx=5, pady=5)


        apply_filters_button.bind(
            "<Enter>",
            lambda e: apply_filters_button.config(fg="white", bg="#8B8989"))

        apply_filters_button.bind(
            "<Leave>",
            lambda e: apply_filters_button.config(fg="black", bg="#CDC9C9"))


        # ------------------ CLEAR FILTER BUTTON ------------------




        clear_filters_button = tk.Button(
            filters_frame,
            text="Clear Filters",
            command=clear_filters,
            bg="#CDC9C9",
            fg="black",
            font=("Helvetica", 12))

        clear_filters_button.grid(row=1, column=11, padx=5, pady=5)


        clear_filters_button.bind(
            "<Enter>",
            lambda e:clear_filters_button.config(fg="white", bg="#8B8989"))

        clear_filters_button.bind(
            "<Leave>",
            lambda e: clear_filters_button.config(fg="black", bg="#CDC9C9"))



        utilities_frame = tk.Frame(movements_window, pady=10)

        utilities_frame.pack()


        total_records_label = tk.Label(
            utilities_frame,
            text="Total records: 0",
            font=("Helvetica", 12))

        total_records_label.pack(side="left", padx=10)


        # ----- EXPORT SELECTED DATA TO EXCEL ------


        export_button = tk.Button(
            utilities_frame,
            text="Export to Excel",
            command=export_movements,
            bg="#CDC9C9",
            fg="black",
            font=("helvetica", 12))

        export_button.pack(side="right", padx=10)



        export_button.bind(
            "<Enter>",
            lambda e:export_button.config(fg="white", bg="#8B8989"))

        export_button.bind(
            "<Leave>",
            lambda e: export_button.config(fg="black", bg="#CDC9C9"))



        # ---------- GENERATE PDF INVENTORY REPORT ----------


        inventory_button = tk.Button(
            utilities_frame,
            text="Generate PDF Inventory",
            command=generate_inventory,
            bg="#2a9d8f",
            fg="white",
            font=("Helvetica", 12))

        inventory_button.pack(side="right", padx=10)


        inventory_button.bind(
            "<Enter>",
            lambda e:inventory_button.config(fg="white"))

        inventory_button.bind(
            "<Leave>",
            lambda e: inventory_button.config(fg="black"))




        # ========== MOVEMENT TREEVIEW ============

        movements_tree = ttk.Treeview(movements_window,
                                columns=("code", "product", "type", "quantity", "date", "location"),
                                show="headings")


        movements_tree.heading("code", text="Code")

        movements_tree.heading("product", text="Product")

        movements_tree.heading("type", text="Type")

        movements_tree.heading("quantity", text="Quantity")

        movements_tree.heading("location", text="Location")

        movements_tree.heading("date", text="Date")




        movements_tree.column("code", width=200, anchor="center")

        movements_tree.column("product", width=200, anchor="center")

        movements_tree.column("type", width=100, anchor="center")

        movements_tree.column("quantity", width=100, anchor="center")

        movements_tree.column("location", width=150, anchor="center")

        movements_tree.column("date", width=200, anchor="center")



        movements_tree.pack(fill="both", expand=True)

        load_movements()























    # ------------------------- MAIN WINDOW BUTTON FRAME -----------------------------

    button_frame = tk.Frame(main_window, pady=10)

    button_frame.pack(fill="x")


    add_product_button = tk.Button(button_frame,
                              text="Add Product",
                              command=add_product,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=20)


    add_product_button.pack(side="left", padx=(90 , 5))



    add_product_button.bind(
    "<Enter>",
    lambda e: add_product_button.config(fg="white", bg="#8B8989"))



    add_product_button.bind(
    "<Leave>",
    lambda e: add_product_button.config(fg="black", bg="#CDC9C9"))





    remove_stock_button = tk.Button(button_frame,
                              text="Remove Product Stock",
                              command=remove_product_stock,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=20)


    remove_stock_button.pack(side="left", padx=10)


    remove_stock_button.bind(
    "<Enter>",
    lambda e: remove_stock_button.config(fg="white", bg="#8B8989"))



    remove_stock_button.bind(
    "<Leave>",
    lambda e: remove_stock_button.config(fg="black", bg="#CDC9C9"))




    movements_button = tk.Button(button_frame,
                              text="Movement History",
                              command=show_movements,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=25)


    movements_button.pack(side="left", padx=10)


    movements_button.bind(
    "<Enter>",
    lambda e: movements_button.config(fg="white", bg="#8B8989"))



    movements_button.bind(
    "<Leave>",
    lambda e: movements_button.config(fg="black", bg="#CDC9C9"))




    manage_users_button = tk.Button(button_frame,
                              text="Manage Users",
                              command=open_user_management,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=20)


    manage_users_button.pack(side="left", padx=10)


    manage_users_button.bind(
    "<Enter>",
    lambda e: manage_users_button.config(fg="white", bg="#8B8989"))



    manage_users_button.bind(
    "<Leave>",
    lambda e: manage_users_button.config(fg="black", bg="#CDC9C9"))



    # ------------------------- FILTER FRAME -----------------------------

    filter_frame = tk.Frame(main_window, pady=10)

    filter_frame.pack(fill="x")



    tk.Label(filter_frame,
             text="Filter by product or location",
             font=("Helvetica", 12)).pack(side="left", padx=10)



    global filter_entry


    filter_entry = tk.Entry(filter_frame,
                              font=("Helvetica", 12),
                              width=50)

    filter_entry.pack(side="left", padx=10)




    apply_filter_button = tk.Button(filter_frame,
                            text="Apply Filter",
                            command=filter_product_or_location,
                            bg="#CDC9C9",
                            fg="black",
                            width=15,
                            font=("Helvetica", 12))

    apply_filter_button.pack(side="left", padx=10, pady=10)

    apply_filter_button.bind(
    "<Enter>",
    lambda e: apply_filter_button.config(fg="white", bg="#8B8989"))



    apply_filter_button.bind(
    "<Leave>",
    lambda e: apply_filter_button.config(fg="black", bg="#CDC9C9"))




    clear_filter_button = tk.Button(filter_frame,
                            text="Clear Filter",
                            command=lambda: [filter_entry.delete(0, tk.END), load_data()],
                            bg="#CDC9C9",
                            fg="black",
                            width=15,
                            font=("Helvetica", 12))

    clear_filter_button.pack(side="left", padx=10, pady=10)

    clear_filter_button.bind(
    "<Enter>",
    lambda e: clear_filter_button.config(fg="white", bg="#8B8989"))



    clear_filter_button.bind(
    "<Leave>",
    lambda e: clear_filter_button.config(fg="black", bg="#CDC9C9"))


    # ------------------------- LIST FRAME -----------------------------

    list_frame = tk.Frame(main_window, pady=20)

    list_frame.pack(fill="both", expand=True)

    global products_tree


    products_tree = ttk.Treeview(list_frame,
                                 columns=("Code", "Name", "Price", "Quantity", "Location", "ID"),
                                 show="headings")




    vertical_scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=products_tree.yview)



    products_tree.configure(yscrollcommand=vertical_scrollbar.set)


    vertical_scrollbar.pack(side="right", fill="y")



    products_tree.heading("Code", text="Code")

    products_tree.heading("Name", text="Name")

    products_tree.heading("Price", text="Price")

    products_tree.heading("Quantity", text="Quantity")

    products_tree.heading("Location", text="Location")

    products_tree.heading("ID", text="ID")





    products_tree.column("Code", width=200, anchor="center")

    products_tree.column("Name", width=200, anchor="center")

    products_tree.column("Price", width=100, anchor="center")

    products_tree.column("Quantity", width=100, anchor="center")

    products_tree.column("Location", width=150, anchor="center")

    products_tree.column("ID", width=100, anchor="center")


    products_tree.pack(fill="both", expand=True)


    products_tree.bind('<<TreeviewSelect>>', populate_product_fields)


    load_data()



    main_window.mainloop()




























# ====== LOGIN PAGE INTERFACE ======

login_window = tk.Tk()

login_window.title("Login - Inventory Management")

login_window.state('zoomed')

login_window.resizable(False, False)



def absolute_path(relative_path):

    if hasattr(sys, '_MEIPASS'):

        return os.path.join(sys._MEIPASS, relative_path)

    return os.path.join(os.path.abspath("."), relative_path)


icon = PhotoImage(file=absolute_path("logo.png"))

login_window.iconphoto(True, icon)





# ------ BACKGROUND IMAGE ------

original_image = Image.open(absolute_path("background.png"))


background_label = tk.Label(login_window)

background_label.place(x=0, y=0, relwidth=1, relheight=1)



def update_background(event=None):

    window_width = login_window.winfo_width()

    window_height = login_window.winfo_height()


    if window_width < 10 or window_height < 10:
        return


    image_width, image_height = original_image.size


    scale = max(window_width / image_width, window_height / image_height)

    new_width = int(image_width * scale)

    new_height = int(image_height * scale)



    resized_image = original_image.resize(
        (new_width, new_height),
        Image.LANCZOS

    )



    left = (new_width - window_width) // 2
    top = (new_height - window_height) // 2
    right = left + window_width
    bottom = top + window_height



    cropped_image = resized_image.crop((left, top, right, bottom))

    overlay = Image.new("RGBA", cropped_image.size, (0, 0, 0, 90))



    final_image = Image.alpha_composite(
        cropped_image.convert("RGBA"),
        overlay

    )

    tk_image = ImageTk.PhotoImage(final_image)

    background_label.config(image=tk_image)

    background_label.image = tk_image


login_window.after(100, update_background)

login_window.bind("<Configure>", update_background)



# ------ CENTERED LOGIN FRAME ------

login_frame = tk.Frame(login_window, padx=40, pady=40, bg="#F8F9FA")

login_frame.place(relx=0.5, rely=0.5, anchor='center')


login_title_label = tk.Label(login_frame, text="Inventory Management", bg="#F8F9FA", font=("Helvetica", 20, "bold"))
login_title_label.grid(row=0, column=0, columnspan=2, pady=20)




tk.Label(login_frame, text="Username", bg="#F8F9FA", font=("Helvetica", 14)).grid(row=1, column=0, pady=10, sticky='e')

username_entry = tk.Entry(login_frame, font=("Helvetica", 14))
username_entry.grid(row=1, column=1, pady=10)




tk.Label(login_frame, text="Password", bg="#F8F9FA", font=("Helvetica", 14)).grid(row=2, column=0, pady=10, sticky='e')

password_entry = tk.Entry(login_frame, font=("Helvetica", 14))
password_entry.grid(row=2, column=1, pady=10)











login_button = tk.Button(login_frame,
                      text="Login",
                      command=verify_login,
                      bg="#2A9D8F",
                      fg="white",
                      font=("Helvetica", 14),
                      width=15)

login_button.grid(row=3, column=0, columnspan=2, pady=10)









register_user_button = tk.Button(login_frame,
                      text="Register User",
                      command=open_user_management,
                      bg="#264653",
                      fg="white",
                      font=("Helvetica", 14),
                      width=15)

register_user_button.grid(row=4, column=0, columnspan=2, pady=10)



create_admin_user()









login_window.mainloop()
