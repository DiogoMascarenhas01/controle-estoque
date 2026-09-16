import tkinter as tk

from tkinter import PhotoImage

from tkinter import ttk, messagebox

from tkcalendar import DateEntry

from pymongo import MongoClient

from bson.objectid import ObjectId

import pandas as pd

import datetime

import bcrypt

import sys

import os 

from PIL import Image, ImageTk

import webbrowser


from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from reportlab.platypus.paragraph import ParagraphStyle

from reportlab.platypus.flowables import KeepTogether

from reportlab.lib import colors

from reportlab.lib.styles import getSampleStyleSheet

from reportlab.lib.pagesizes import A4

from reportlab.pdfgen import canvas

from reportlab.platypus import Image as RLImage

from openpyxl import Workbook

from openpyxl.styles import Alignment, Font



# ====== CONEXÃO COM O BANCO DE DADOS ======


cliente = MongoClient("mongodb://localhost:27017/")


bancoDeDados = cliente["ControleDeEstoque"]


conexao_produtos = bancoDeDados["produtos"]

conexao_movimentacao = bancoDeDados["movimentacao"]

conexao_usuarios = bancoDeDados["usuarios"]




# ====== CRIAÇÃO DO USUÁRIOS ADMINISTRADOR ======

def criar_usuario_admin():
    if conexao_usuarios.count_documents({"username":"a"}) == 0:

        senha = "a"

        hashed = bcrypt.hashpw(senha.encode('utf-8'), bcrypt.gensalt())


        conexao_usuarios.insert_one({
            "username": "a",
            "password": hashed,
            "plain_password": senha
        })


# ====== VERIFICAÇÃO DE LOGIN ======

def verificar_login():

    username = entrada_usuario.get()

    senha = entrada_senha.get()



    if not username or not senha:

        messagebox.showerror("Erro", "Por favor, preencha todos os campos")

        return
    

    usuario = conexao_usuarios.find_one({"username":username})


    if usuario and bcrypt.checkpw(senha.encode('utf-8'), usuario["password"]):

        janela_login.destroy()



        abrir_janela_principal()

    
    else:

        messagebox.showerror("Erro", "Usuário ou senha incorretos")






# =============== AQUI DESENVOLVEREMOS A TELA DE USUÁRIOS =====================


def abrir_tela_usuarios():
    
    janela_usuario = tk.Toplevel()

    janela_usuario.title("Gerenciar Usuários")

    janela_usuario.geometry("800x600")

    janela_usuario.resizable(False, False)

    janela_usuario.update_idletasks()

    largura  = janela_usuario.winfo_width()

    altura = janela_usuario.winfo_height()


    pos_x = (janela_usuario.winfo_screenwidth() // 2 ) - (largura // 2 )

    pos_y = (janela_usuario.winfo_screenheight() // 2 ) - (altura // 2 )


    janela_usuario.geometry(f"{largura}x{altura}+{pos_x}+{pos_y}")



    #====== FUNÇÃO RESPONSÁVEL POR CARREGAROS DADOS DOS USUÁRIOS NO TREEVIEW =====

    def carregar_usuarios():

        for item in tree_usuarios.get_children():

            tree_usuarios.delete(item)


        registros = conexao_usuarios.find()


        for doc in registros:

            tree_usuarios.insert(
                "",
                "end",
                values=(
                    str(doc["_id"]),
                    doc["username"],
                    doc.get("plain_password", "")

                )

            )


    #====== FRAME DO TOPO DA JANELA DE USUÁRIOS =====
    frame_topo = tk.Frame(janela_usuario, pady=10)

    frame_topo.pack()



    tk.Label(
        frame_topo,
        text="Gerenciar Usuários",
        font=("Helvetica", 16, "bold")).pack()
    




    frame_form = tk.Frame(janela_usuario, padx=20, pady=10)

    frame_form.pack()

    tk.Label(
        frame_form,
        text="Usuário",
        font=("Helvetica", 12)
    ).grid(row=0, column=0, sticky="e")


    entrada_novo_usuario = tk.Entry(
        frame_form,
        font=("Helvetica", 12))
    
    entrada_novo_usuario.grid(row=0, column=1, padx=10, pady=5 )




    tk.Label(
        frame_form,
        text="Senha",
        font=("Helvetica", 12)
    ).grid(row=1, column=0, sticky="e")


    entrada_nova_senha = tk.Entry(
        frame_form,
        font=("Helvetica", 12))
    
    entrada_nova_senha.grid(row=1, column=1, padx=10, pady=5 )


    #========= FRAME PARA OS BOTÕES =========

    frame_botoes = tk.Frame(janela_usuario, pady=10)

    frame_botoes.pack()


    #========= BOTÃO ADICIONAR USUÁRIO =========

    def adicionar_usuario():

        username = entrada_novo_usuario.get()

        senha = entrada_nova_senha.get()


        if not username or not senha:

            messagebox.showerror("Erro", "Por favor, preencha todos os campo!!!")

            return
        

        if conexao_usuarios.find_one({"username": username}):

            messagebox.showerror("Erro", "O usuário já existe no banco de dados!!!")

            return
        

        hashed = bcrypt.hashpw(senha.encode('utf-8'), bcrypt.gensalt())


        conexao_usuarios.insert_one({
            "username": username,
            "password": hashed,
            "plain_password": senha
        })

        messagebox.showinfo("Sucesso", "Usuário caastrado com sucesso!!!")


        entrada_novo_usuario.delete(0, tk.END)

        entrada_nova_senha.delete(0, tk.END)


        carregar_usuarios()


        janela_usuario.destroy()




    btn_adicionar_usuario = tk.Button(
        frame_botoes,
        text="Adicionar Usuário",
        command=adicionar_usuario,
        bg="#CDC9C9",
        fg="black",
        font=("Helvetica", 12),
        width=15
    )

    btn_adicionar_usuario.pack(side="left", padx=5)



    btn_adicionar_usuario.bind(
    "<Enter>",
    lambda e: btn_adicionar_usuario.config(fg="white", bg="#8B8989"))  



    btn_adicionar_usuario.bind(
    "<Leave>",
    lambda e: btn_adicionar_usuario.config(fg="black", bg="#CDC9C9"))


    
    #========= BOTÃO EDITAR USUÁRIO =========

    def editar_usuario():

        selecionado = tree_usuarios.selection()

        if not selecionado:

            messagebox.showerror("Erro", "Selecione um usuário")

            return
        

        item = tree_usuarios.item(selecionado)

        doc_id = item["values"][0]


        username = entrada_novo_usuario.get()

        nova_senha = entrada_nova_senha.get()



        if not username or not nova_senha:

            messagebox.showerror("Erro", "Por favor, preencha todo os campos!!!")

            return
        
        hashed = bcrypt.hashpw(nova_senha.encode('utf-8'), bcrypt.gensalt())


        conexao_usuarios.update_one(
            {"_id": ObjectId(doc_id)},
            {"$set":{
                "username": username,
                "password" : hashed,
                "plain_password": nova_senha
            }}
        )


        messagebox.showinfo("Sucesso", f"Usuário '{username}' atualizado com sucesso!!!")


        entrada_novo_usuario.delete(0, tk.END)

        entrada_nova_senha.delete(0, tk.END)


        carregar_usuarios()
        


    btn_editar_usuario = tk.Button(
        frame_botoes,
        text="Editar Usuário",
        command=editar_usuario,
        bg="#CDC9C9",
        fg="black",
        font=("Helvetica", 12),
        width=15
    )

    btn_editar_usuario.pack(side="left", padx=5)
    


    btn_editar_usuario.bind(
    "<Enter>",
    lambda e: btn_editar_usuario.config(fg="white", bg="#8B8989"))  



    btn_editar_usuario.bind(
    "<Leave>",
    lambda e: btn_editar_usuario.config(fg="black", bg="#CDC9C9"))



    #========= BOTÃO EXCLUIR USUÁRIO =========

    def excluir_usuario():

        selecionado = tree_usuarios.selection()

        if not selecionado:

            messagebox.showerror("Erro", "Selecione um usuário para excluir!!!")


            return
        

        item = tree_usuarios.item(selecionado)

        username = item["values"][1]

        doc_id = item["values"][0]


        if messagebox.askyesno("Confirmação", f"Tem certeza que deseja excluir o usuário '{username}?"):

            conexao_usuarios.delete_one({"_id": ObjectId(doc_id)})


            messagebox.showinfo("Sucesso", f"Usuário: '{username}' excluído com sucesso!!!")

            carregar_usuarios()
  

    btn_excluir_usuario = tk.Button(
        frame_botoes,
        text="Excluir Usuário",
        command=excluir_usuario, 
        bg="#CDC9C9",
        fg="black",
        font=("Helvetica", 12),
        width=15
    )

    btn_excluir_usuario.pack(side="left", padx=5)
    


    btn_excluir_usuario.bind(
    "<Enter>",
    lambda e: btn_excluir_usuario.config(fg="white", bg="#8B8989"))  



    btn_excluir_usuario.bind(
    "<Leave>",
    lambda e: btn_excluir_usuario.config(fg="black", bg="#CDC9C9"))


    #========= TREEVIEW PARA EXIBIÇÃO DA TABELA DE USUÁRIOS =========

    def preencher_campos_usuario(event):

        selecionado = tree_usuarios.selection()

        if selecionado:

            item = tree_usuarios.item(selecionado)

            valores = item['values']


            entrada_novo_usuario.delete(0, tk.END)

            entrada_novo_usuario.insert(0, valores[1])


            entrada_nova_senha.delete(0, tk.END)

            entrada_nova_senha.insert(0, valores[2])



    frame_lista = tk.Frame(janela_usuario, pady=10)

    frame_lista.pack(fill="both", expand=True)


        
    tree_usuarios = ttk.Treeview(
        frame_lista,
        columns=("ID", "Username", "Password"),
        show="headings"
    )


    tree_usuarios.heading("ID", text="ID")

    tree_usuarios.column("ID", width=100, anchor="center")




    tree_usuarios.heading("Username", text="Usuário")

    tree_usuarios.column("Username", width=200, anchor="center")




    tree_usuarios.heading("Password", text="Senha")

    tree_usuarios.column("Password", width=200, anchor="center")



    tree_usuarios.pack(fill="both", expand=True)


    tree_usuarios.bind("<<TreeviewSelect>>", preencher_campos_usuario)

    carregar_usuarios()

        










# FUNÇÃO RESPONSÁVEL POR DAR BAIXA NO PRODUTO

def dar_baixa_produto():

    selecionado = tree_produtos.selection()

    if not selecionado:

        messagebox.showerror("Erro", "Selecione um produto para fazer a retirada")

        return
    


    item = tree_produtos.item(selecionado)

    doc_id = item["values"][5]

    codigo_produto = item["values"][0]

    nome_produto = item["values"][1]

    quantidade_atual = item["values"][3]

    quantidade_baixa = entrada_quantidade.get()


    if not quantidade_baixa:

        messagebox.showerror("Erro", "Por favor, informe a quantidade a ser retirada!!!")

        return
    
    try:

        quantidade_baixa = int(quantidade_baixa)

    except ValueError:

        messagebox.showerror("Erro", "A quantidade a ser retirada deve ser número!!!")

        return
    
    if quantidade_baixa > quantidade_atual:

        messagebox.showerror("Erro", "Quantidade superior ao que tem no estoque!!!")

        return
    
    nova_quantidade = quantidade_atual - quantidade_baixa


    conexao_produtos.update_one(
        {"_id": ObjectId(doc_id)},
        {"$set": {"quantidade": nova_quantidade}}
        )
    

    registrar_movimentacao(codigo_produto, nome_produto, "SAÍDA", quantidade_baixa)

    messagebox.showinfo("Sucesso", f"A retirada de {quantidade_baixa} unidades foi realizada com sucesso!!!")


    carregar_dados()


    limpar_campos()
    

    














#FUNÇÃO RESPONSÁVEL POR CADASTRAR UM NOVO PRODUTO NO SISTEMA

def adicionar_produto():
    
    codigo = entrada_codigo.get().strip().upper()

    nome = entrada_nome.get().strip().upper()

    preco = entrada_preco.get().strip().upper()

    quantidade = entrada_quantidade.get().strip().upper()

    localizacao = entrada_localizacao.get().strip().upper()



    if not codigo or not  nome or not preco or not quantidade or not localizacao:

        messagebox.showerror("Erro", "Todos os campos são obrigatórios")

        return
    

    try:

        preco = float(preco)

        quantidade = int(quantidade)

    
    except ValueError:

        messagebox.showerror("Erro", "Para inserir preço e quantidade é necessários digitar números e não letras!!!")

        return
    

    conexao_produtos.insert_one({
        "codigo": codigo,
        "nome" : nome,
        "preco": preco,
        "quantidade":quantidade,
        "localizacao": localizacao

    })



    registrar_movimentacao(codigo, nome, "ENTRADA", quantidade)

    messagebox.showinfo("Sucesso", "Parabéns, o produto foi cadastrado con sucesso!!!")



    carregar_dados()

    limpar_campos()




def registrar_movimentacao(codigo, produto, tipo, quantidade):

    data_atual = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    produto_encontrado = conexao_produtos.find_one({"codigo": codigo})

    conexao_movimentacao.insert_one({

        "codigo": codigo,
        "produto": produto,
        "tipo":tipo,
        "quantidade":quantidade,
        "data": data_atual,
        "localizacao": produto_encontrado.get("localizacao", "")

    })


def limpar_campos():

    entrada_codigo.delete(0, tk.END)

    entrada_nome.delete(0, tk.END)

    entrada_preco.delete(0, tk.END)

    entrada_quantidade.delete(0, tk.END)

    entrada_localizacao.delete(0, tk.END)



# RESPONSÁVEL POR CARREGAR OS DADOS NO TREEVIEW

def carregar_dados():
    
    for item in tree_produtos.get_children():

        tree_produtos.delete(item)



    registros = conexao_produtos.find({"quantidade": {"$gt": 0}})

    for doc in registros:

        quantidade = doc.get("quantidade", 0)


        item_id = tree_produtos.insert(
            "",
            "end",
            values=(
                doc.get("codigo", ""),
                doc.get("nome", ""),
                doc.get("preco", ""),
                quantidade,
                doc.get("localizacao", ""),
                str(doc.get("_id", ""))
            )
        )



# FUNÇÃO RESPONSÁVEL POR FILTRAR PRODUTO OU LOCALIZAÇÃO
def filtrar_produto_localizacao():

    termo = entrada_filtro.get().strip()

    if not termo:

        carregar_dados()

        return
    
    for item in tree_produtos.get_children():

        tree_produtos.delete(item)


    regex = {"$regex" : termo, "$options" : "i"}


    try:

        preco_valor = float(termo)

        quantidade_valor = int(termo)

    
    except ValueError:

        preco_valor = None

        quantidade_valor = None
    
    

    query = {"$or": [
        {"nome" : regex},
        {"localizacao" : regex}
    ]}


    if preco_valor is not None:

        query["$or"].append({"preco": preco_valor})

    
    if quantidade_valor is not None:

        query["$or"].append({"quantidade": quantidade_valor})


    
    registros = conexao_produtos.find(query)


    for doc in registros:

        quantidade = doc["quantidade"]

        item_id = tree_produtos.insert(
            "",
            "end",
            values=(
                doc.get("codigo", ""),
                doc.get("nome", ""),
                doc.get("preco", ""),
                quantidade,
                doc.get("localizacao", ""),
                str(doc.get("_id", ""))

            )

        )













# FUNÇÃO PARA REALIZAR O PRREENCIMENTO DOS CAMPOS COM OS DADOS SELECIONADOS PELO USUÁRIO


def preencher_campos_produto(event):

    selecionado = tree_produtos.selection()

    if selecionado:

        item = tree_produtos.item(selecionado)

        valores = item['values']


        entrada_codigo.delete(0, tk.END)
        entrada_codigo.insert(0, valores[0])


        entrada_nome.delete(0, tk.END)
        entrada_nome.insert(0, valores[1])


        entrada_preco.delete(0, tk.END)
        entrada_preco.insert(0, valores[2])



        entrada_quantidade.delete(0, tk.END)
        entrada_quantidade.insert(0, valores[3])



        entrada_localizacao.delete(0, tk.END)
        entrada_localizacao.insert(0, valores[4])


        entrada_quantidade.focus_set()



    
    







#======= FUNÇÕES DE FORMA PROVISÓRIA =====



def exibir_movimentacoes():
    print("Aguardando a função exibir_movimentacoes")

    




def excluir_usuario():
    print("Aguardando a função excluir_usuario")



#^^^^^^ FUNÇÕES DE FORMA PROVISÓRIA ^^^^




# ====== JANELA PRINCIPAL DO SISTEMA ======

def  abrir_janela_principal():

    global janela_principal 

    janela_principal = tk.Tk()

    janela_principal.title("Sistema de controle de estoque")

    janela_principal.state('zoomed')

    janela_principal.resizable(True, True)



    # ------------------------ FRAME TOPO ----------------------------


    frame_topo = tk.Frame(janela_principal, bg="#E8E8E8", height=70)

    frame_topo.pack(fill="x")


    lbl_titulo = tk.Label(frame_topo, 
                          text="Sistema de controle de estoque", 
                          bg="#E8E8E8", 
                          fg="black",
                          font=("Helvetica", 20, "bold"))
    

    lbl_titulo.pack(pady=10)




    # ------------------------- FRAME DADOS -----------------------------


    frame_dados = tk.Frame(janela_principal, padx=20, pady=10)

    frame_dados.pack(fill="x")

    # -------------------------------------------------------------------

    tk.Label(frame_dados,
             text="Código",
             font=("Helvetica", 12)).grid(row=0, column=0, sticky="e")
    

    global entrada_codigo


    entrada_codigo = tk.Entry(frame_dados, font=("Helvetica", 12))

    entrada_codigo.grid(row=0, column=1, padx=10, pady=5)

    # -------------------------------------------------------------------

    tk.Label(frame_dados,
             text="Nome",
             font=("Helvetica", 12)).grid(row=0, column=2, sticky="e")
    
    global entrada_nome

    entrada_nome = tk.Entry(frame_dados, font=("Helvetica", 12))

    entrada_nome.grid(row=0, column=3, padx=10, pady=5)


    # -------------------------------------------------------------------

    tk.Label(frame_dados,
             text="Preço",
             font=("Helvetica", 12)).grid(row=0, column=4, sticky="e")
    
    global entrada_preco

    entrada_preco = tk.Entry(frame_dados, font=("Helvetica", 12))

    entrada_preco.grid(row=0, column=5, padx=10, pady=5)


    # -------------------------------------------------------------------

    tk.Label(frame_dados,
             text="Quantidade",
             font=("Helvetica", 12)).grid(row=0, column=6, sticky="e")
    
    global entrada_quantidade

    entrada_quantidade = tk.Entry(frame_dados, font=("Helvetica", 12))

    entrada_quantidade.grid(row=0, column=7, padx=10, pady=5)


    # -------------------------------------------------------------------

    tk.Label(frame_dados,
             text="Localização",
             font=("Helvetica", 12)).grid(row=0, column=8, sticky="e")
    
    global entrada_localizacao

    entrada_localizacao = tk.Entry(frame_dados, font=("Helvetica", 12))

    entrada_localizacao.grid(row=0, column=9, padx=10, pady=5)



    #========================== JANELA DO BOTÃO DE HISTORICO DE MOVIMENTAÇÕES ==========================
    def exibir_movimentacoes():

        janela_movimentacoes = tk.Toplevel(janela_principal)

        janela_movimentacoes.title("Histórico de Movimentações")

        janela_movimentacoes.resizable(False, False)

        janela_movimentacoes.update_idletasks()

        largura = 1500

        altura = 600

        pos_x = (janela_movimentacoes.winfo_screenwidth() // 2) - (largura // 2)

        pos_y = (janela_movimentacoes.winfo_screenheight() // 2) - (altura // 2)


        janela_movimentacoes.geometry(f"{largura}x{altura}+{pos_x}+{pos_y}")



        # ------------------------- FRAME DA JANELA HISTÓRICO DE MOVIMENTAÇÕES ( FRAME_FILTROS) -----------------------------
        frame_filtros = tk.Frame(janela_movimentacoes, pady=10)

        frame_filtros.pack()



        tk.Label(frame_filtros, text="Histórico de Movimentações",
                 font=("arial", 22, "bold")).grid(row=0, column=0, columnspan=8, padx=(500, 0), pady=(10, 30))
        


        tk.Label(frame_filtros, text="Código",
                 font=("arial", 12)).grid(row=1, column=0, padx=5)
        

        entrada_filtro_codigo = tk.Entry(frame_filtros,
                                         font=("Helvetica", 12))
                                        
        entrada_filtro_codigo.grid(row=1, column=1, padx=5)






        tk.Label(frame_filtros, text="Produto",
                 font=("arial", 12)).grid(row=1, column=2, padx=5)
        

        entrada_filtro_produto = tk.Entry(frame_filtros,
                                         font=("Helvetica", 12))
                                        
        entrada_filtro_produto.grid(row=1, column=3, padx=5)





        tk.Label(frame_filtros, text="Tipo",
                 font=("arial", 12)).grid(row=1, column=4, padx=5)
        

        entrada_filtro_tipo = tk.Entry(frame_filtros,
                                         font=("Helvetica", 12))
                                        
        entrada_filtro_tipo.grid(row=1, column=5, padx=5)





        tk.Label(frame_filtros, text="Data Inicial",
                 font=("arial", 12)).grid(row=1, column=6, padx=5)
        

        entrada_data_inicial = DateEntry(frame_filtros,
                                         font=("Helvetica", 12),
                                         date_pattern = 'yyyy-MM-dd')
                                        
        entrada_data_inicial.grid(row=1, column=7, padx=5)


        # ------------ BUSCA AUTOMATICAMENTE PELA MOVIMENTAÇÃO MAIS ANTIGA ------------

        primeira_movimentacao = conexao_movimentacao.find_one(
            {},
            sort=[("data", 1)]
        )


        if primeira_movimentacao:

            data_primeira = datetime.datetime.strptime(
                primeira_movimentacao["data"],
                "%Y-%m-%d %H:%M:%S"
            ).date()

            entrada_data_inicial.set_date(data_primeira)





        tk.Label(frame_filtros,
                 text="Data Final:",
                 font=("arial", 12)).grid(row=1, column=8, padx=5)
        

        entrada_data_final = DateEntry(frame_filtros,
                                         font=("Helvetica", 12),
                                         date_pattern = 'yyyy-MM-dd')
                                        
        entrada_data_final.grid(row=1, column=9, padx=5)


 
        #------------------ BOTÃO PARA EXECUTAR O FILTRO ------------------

        def carregar_movimentacoes():

            for item in tree_mov.get_children():

                tree_mov.delete(item)

            
            filtro = {}


            codigo = entrada_filtro_codigo.get().strip()

            produto = entrada_filtro_produto.get().strip()

            tipo = entrada_filtro_tipo.get().strip().upper()

            data_inicial = entrada_data_inicial.get_date()

            data_final = entrada_data_final.get_date()


            if codigo:

                filtro["codigo"] = {"$regex": codigo, "$options": "i"}

            
            if produto:

                filtro["produto"] = {"$regex": produto, "$options": "i"}


            if tipo:

                filtro["tipo"] = {"$regex": tipo, "$options": "i"}



            if data_inicial or data_final:

                filtro["data"] = {}

                if data_inicial:

                    filtro["data"]["$gte"] = data_inicial.strftime("%Y-%m-%d")


                if data_final:

                    filtro["data"]["$lte"] = data_final.strftime("%Y-%m-%d") + "23:59:59"


            registros = conexao_movimentacao.find(filtro).sort("data", 1)

            total = 0


            for doc in registros:

                tree_mov.insert("", "end",
                                values=(
                                    doc.get("codigo", ""),
                                    doc.get("produto", ""),
                                    doc.get("tipo", ""),
                                    doc.get("quantidade", ""),
                                    doc.get("data", ""),
                                    doc.get("localizacao", "")))
                
                total += 1

            lbl_total_registros.config(text=f"Total de registros: {total}")





        def limpar_filtros():

            entrada_filtro_codigo.delete(0, tk.END)

            entrada_filtro_produto.delete(0, tk.END)

            entrada_filtro_tipo.delete(0, tk.END)


            carregar_movimentacoes()





        # FUNÇÃO RESPONSÁVEL POR EXPORTAR O HISTÓRICO DE MOVIMENTAÇÕES PARA UM ARQUIVO EXCEL
        def exportar_movimentacoes():

            filtro = {}


            codigo = entrada_filtro_codigo.get().strip()

            produto = entrada_filtro_produto.get().strip()

            tipo = entrada_filtro_tipo.get().strip()



            data_inicial = entrada_data_inicial.get_date()

            data_final = entrada_data_final.get_date()


            if codigo:
                filtro["codigo"] = {"$regex": codigo, "$options": "i"}


            if produto:
                filtro["produto"] = {"$regex": produto, "$options": "i"}


            if tipo:
                filtro["tipo"] = {"$regex": tipo, "$options": "i"}




            if data_inicial or data_final:

                filtro["data"] = {}


                if data_inicial:

                    filtro["data"]["$gte"] = data_inicial.strftime("%Y-%m-%d")

                if data_final:

                    filtro["data"]["$lte"] = data_final.strftime("%Y-%m-%d") + "23:59:59"


            registros = list(conexao_movimentacao.find(filtro).sort("data", 1))


            if not registros:

                messagebox.showinfo("Exportar", "Não há registros para exportar")

                return
            

            wb = Workbook()

            ws = wb.active

            ws.title="Movimentações"

            cabecalho = ["CÓDIGO", "PRODUTO", "TIPO", "QUANTIDADE", "LOCALIZAÇÃO", "DATA"]

            ws.append(cabecalho)

            for col in ws[1]:

                col.font= Font(bold=True)

                col.alignment = Alignment(horizontal="center", vertical="center")

            
            for doc in registros:

                linha = [
                    str(doc.get("codigo", "")).upper(),
                    str(doc.get("produto", "")).upper(),
                    str(doc.get("tipo", "")).upper(),
                    str(doc.get("quantidade", "")).upper(),
                    str(doc.get("localizacao", "")).upper(),
                    str(doc.get("data", "")).upper()
                ]

                ws.append(linha)


            for row in ws.iter_rows():

                for cell in row:

                    cell.alignment = Alignment(horizontal="center", vertical="center")


            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 30

            
            arquivo = "movimentacoes_filtradas.xlsx"

            wb.save(arquivo)


            messagebox.showinfo("Exportar", f"Dados exportados com sucesso para: {arquivo}!!! ")




        # FUNÇÃO PARA FORMATAR VALORES PARA PADRÃO BRASILEIRO
        def formatar_moeda_br(valor):
            return f"R${valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        


        # FUNÇÃO PARA GERAR INVENTÁRIO
        def gerar_inventario():

            filtro = {}

            codigo = entrada_filtro_codigo.get().strip()

            produto = entrada_filtro_produto.get().strip()

            tipo = entrada_filtro_tipo.get().strip()



            if codigo:
                filtro["codigo"] = {"$regex": codigo, "$options": "i"}


            if produto:
                filtro["produto"] = {"$regex": produto, "$options": "i"}


            if tipo:
                filtro["tipo"] = {"$regex": tipo, "$options": "i"}




            registros = list(
                conexao_produtos.find(filtro).sort("codigo", 1)

            )


            if not registros:

                messagebox.showinfo("Inventário", "Nenhum produto encontrado para gerar o inventário!!!")
                return
            

            pasta = "INVENTÁRIOS"

            os.makedirs(pasta, exist_ok=True)

            data_hora = datetime.datetime.now().strftime("%d-%m-%Y_%H-%M-%S")


            nome_arquivo = f"inventario_{data_hora}.pdf"


            arquivo_pdf = os.path.join(pasta, nome_arquivo)



            doc = SimpleDocTemplate(
                arquivo_pdf,
                pagesize=A4,
                topMargin=25,
                bottomMargin=25,
                leftMargin=30,
                rightMargin=30
            )


            elementos = []

            estilos = getSampleStyleSheet()



            estilo_titulo = ParagraphStyle(
                "TituloERP",
                parent=estilos["title"],
                alignment=0,
                spaceBefore=0,
                spaceAfter=6,
                leftIndent=40
            )


            estilo_subtitulo = ParagraphStyle(
                "SubTituloERP",
                parent=estilos["Heading4"],
                alignment=0,
                spaceBefore=0,
                spaceAfter=4,
                leftIndent=20
            )


            def caminho_absoluto(relativo):

                if hasattr(sys, "_MEIPASS"):
                    return os.path.join(sys._MEIPASS, relativo)
                

                return os.path.join(os.path.abspath("."), relativo)
            


            caminho_logo = caminho_absoluto("logo.png")

            logo = None

            if os.path.exists(caminho_logo):

                logo = RLImage(caminho_logo, width=70, height=70)


            titulo = Paragraph(
                "<b>RELATÓRIO DE INVENTÁRIO</b>",
                estilo_titulo
            )

            cabecalho_dados = [[logo, titulo]]

            tabela_cabecalho = Table(
                cabecalho_dados,
                colWidths=[80, 400]
            
            )


            tabela_cabecalho.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))

            elementos.append(tabela_cabecalho)
            elementos.append(Spacer(1, 15))


            nome_relatorio = f"Inventário - {data_hora}"

            elementos.append(
                Paragraph(f"<b>{nome_relatorio}</b>", estilo_subtitulo )
                
            )

            elementos.append(Spacer(1, 10))



            dados = [
                ["Código", "Produto", "Quantidade", "Preço Unit.", "Total"]
            ]


            total_geral = 0



            for produto_db in registros:

                codigo = produto_db.get("codigo", "")
                nome_produto = produto_db.get("nome", "")
                quantidade = produto_db.get("quantidade", 0)
                preco = float(produto_db.get("preco", 0))


                total_item = preco * quantidade


                total_geral += total_item




                dados.append([
                    codigo,
                    nome_produto,
                    quantidade,
                    formatar_moeda_br(preco),
                    formatar_moeda_br(total_item)
                ])


            elementos.append(
                Paragraph(
                    f"<b>VALOR TOTAL DO INVENTÁRIO: {formatar_moeda_br(total_geral)}</b>",
                    estilo_subtitulo
                )
            )

            elementos.append(Spacer(1, 15))


            tabela = Table(dados, colWidths=[70, 180, 80, 90, 90 ])

            tabela.repeatRows = 1

            tabela.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5,  colors.black),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("ALIGN", (2, 1), (-1, -1), "CENTER"),
                ("PADDING", (0, 0), (-1, -1), 6),

            ]))

            elementos.append(tabela)


            doc.build(elementos)

            messagebox.showinfo("Sucesso", "Inventário gerado com sucesso!!!")


            caminho_absoluto_pdf = os.path.abspath(arquivo_pdf)

            webbrowser.open(f"file://{caminho_absoluto_pdf}")


                


                 




      
           
        






            
                               












            



        btn_aplicar_filtro = tk.Button(
            frame_filtros,
            text="Aplicar Filtros",
            command=carregar_movimentacoes,
            bg="#CDC9C9",
            fg="black",
            font=("Helvetica", 12)
        )

        btn_aplicar_filtro.grid(row=1, column=10, padx=5, pady=5)


        btn_aplicar_filtro.bind(
            "<Enter>",
            lambda e: btn_aplicar_filtro.config(fg="white", bg="#8B8989"))
        
        btn_aplicar_filtro.bind(
            "<Leave>",
            lambda e: btn_aplicar_filtro.config(fg="black", bg="#CDC9C9"))
        
        
        #------------------ BOTÃO PARA LIMPAR O FILTRO ------------------
      



        btn_limpar_filtros = tk.Button(
            frame_filtros,
            text="Limpar Filtros",
            command=limpar_filtros,
            bg="#CDC9C9",
            fg="black",
            font=("Helvetica", 12))
        
        btn_limpar_filtros.grid(row=1, column=11, padx=5, pady=5)


        btn_limpar_filtros.bind(
            "<Enter>",
            lambda e:btn_limpar_filtros.config(fg="white", bg="#8B8989"))
        
        btn_limpar_filtros.bind(
            "<Leave>",
            lambda e: btn_limpar_filtros.config(fg="black", bg="#CDC9C9"))
        


        frame_utulitarios = tk.Frame(janela_movimentacoes, pady=10)

        frame_utulitarios.pack()


        lbl_total_registros = tk.Label(
            frame_utulitarios,
            text="Total de registros: 0",
            font=("Helvetica", 12))
        
        lbl_total_registros.pack(side="left", padx=10)


        # ----- BOTÃO PARA EXPORTAR OS DADOS ESCOLHIDOS PARA O EXCEL ------
        

        btn_exportar = tk.Button(
            frame_utulitarios,
            text="Exportar para o Excel",
            command=exportar_movimentacoes,
            bg="#CDC9C9",
            fg="black",
            font=("helvetica", 12))
        
        btn_exportar.pack(side="right", padx=10)



        btn_exportar.bind(
            "<Enter>",
            lambda e:btn_exportar.config(fg="white", bg="#8B8989"))
        
        btn_exportar.bind(
            "<Leave>",
            lambda e: btn_exportar.config(fg="black", bg="#CDC9C9"))
        


        #---------- BOTÃO PARA GERAR O INVENTÁRIOS EM PDF --------------


        btn_inventario = tk.Button(
            frame_utulitarios,
            text="Gerar Inventário em PDF",
            command=gerar_inventario,
            bg="#2a9d8f",
            fg="white",
            font=("Helvetica", 12))
        
        btn_inventario.pack(side="right", padx=10)


        btn_inventario.bind(
            "<Enter>",
            lambda e:btn_inventario.config(fg="white"))
        
        btn_inventario.bind(
            "<Leave>",
            lambda e: btn_inventario.config(fg="black"))
        



        # ========== TREEVIEW PARA EXIBIR AS MOVIMENTAÇÕES ============

        tree_mov = ttk.Treeview(janela_movimentacoes,
                                columns=("codigo", "produto", "tipo", "quantidade", "data", "localizacao"),
                                show="headings")
                                

        tree_mov.heading("codigo", text="Código")

        tree_mov.heading("produto", text="Produto")

        tree_mov.heading("tipo", text="Tipo")

        tree_mov.heading("quantidade", text="Quantidade")

        tree_mov.heading("localizacao", text="Localização")

        tree_mov.heading("data", text="Data")




        tree_mov.column("codigo", width=200, anchor="center")

        tree_mov.column("produto", width=200, anchor="center")

        tree_mov.column("tipo", width=100, anchor="center")

        tree_mov.column("quantidade", width=100, anchor="center")

        tree_mov.column("localizacao", width=150, anchor="center")

        tree_mov.column("data", width=200, anchor="center")



        tree_mov.pack(fill="both", expand=True)

        carregar_movimentacoes()

        

                 




        














    # ------------------------- FRAME BOTOES DA JANELA PRINCIPAL -----------------------------

    frame_botoes = tk.Frame(janela_principal, pady=10)

    frame_botoes.pack(fill="x")


    btn_adicionar = tk.Button(frame_botoes,
                              text="Adicionar Produto",
                              command=adicionar_produto,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=20)
                              
                    
    btn_adicionar.pack(side="left", padx=(90 , 5))



    btn_adicionar.bind(
    "<Enter>",
    lambda e: btn_adicionar.config(fg="white", bg="#8B8989"))  



    btn_adicionar.bind(
    "<Leave>",
    lambda e: btn_adicionar.config(fg="black", bg="#CDC9C9")) 





    btn_baixa = tk.Button(frame_botoes,
                              text="Retirada de Produto",
                              command=dar_baixa_produto,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=20)
                              
                    
    btn_baixa.pack(side="left", padx=10)


    btn_baixa.bind(
    "<Enter>",
    lambda e: btn_baixa.config(fg="white", bg="#8B8989"))  



    btn_baixa.bind(
    "<Leave>",
    lambda e: btn_baixa.config(fg="black", bg="#CDC9C9")) 




    btn_movimentacoes = tk.Button(frame_botoes,
                              text="Histórico de movimentações",
                              command=exibir_movimentacoes,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=25)
                              
                    
    btn_movimentacoes.pack(side="left", padx=10)


    btn_movimentacoes.bind(
    "<Enter>",
    lambda e: btn_movimentacoes.config(fg="white", bg="#8B8989"))  



    btn_movimentacoes.bind(
    "<Leave>",
    lambda e: btn_movimentacoes.config(fg="black", bg="#CDC9C9")) 




    btn_gerenciar_usuarios = tk.Button(frame_botoes,
                              text="Gerenciar Usuários",
                              command=abrir_tela_usuarios,
                              bg="#CDC9C9",
                              fg="black",
                              font=("Helvetica", 12),
                              width=20)
                              
                    
    btn_gerenciar_usuarios.pack(side="left", padx=10)


    btn_gerenciar_usuarios.bind(
    "<Enter>",
    lambda e: btn_gerenciar_usuarios.config(fg="white", bg="#8B8989"))  



    btn_gerenciar_usuarios.bind(
    "<Leave>",
    lambda e: btn_gerenciar_usuarios.config(fg="black", bg="#CDC9C9")) 



    # ------------------------- FRAME FILTRO -----------------------------

    frame_filtro = tk.Frame(janela_principal, pady=10)

    frame_filtro.pack(fill="x")



    tk.Label(frame_filtro,
             text="Filtrar produto ou localização",
             font=("Helvetica", 12)).pack(side="left", padx=10)
    


    global entrada_filtro


    entrada_filtro = tk.Entry(frame_filtro,
                              font=("Helvetica", 12),
                              width=50)

    entrada_filtro.pack(side="left", padx=10)




    btn_filtrar = tk.Button(frame_filtro,
                            text="Aplicar Filtro",
                            command=filtrar_produto_localizacao,
                            bg="#CDC9C9",
                            fg="black",
                            width=15,
                            font=("Helvetica", 12))
    
    btn_filtrar.pack(side="left", padx=10, pady=10)

    btn_filtrar.bind(
    "<Enter>",
    lambda e: btn_filtrar.config(fg="white", bg="#8B8989"))  



    btn_filtrar.bind(
    "<Leave>",
    lambda e: btn_filtrar.config(fg="black", bg="#CDC9C9")) 




    btn_limpar_filtro = tk.Button(frame_filtro,
                            text="Limpar Filtro",
                            command=lambda: [entrada_filtro.delete(0, tk.END), carregar_dados()],
                            bg="#CDC9C9",
                            fg="black",
                            width=15,
                            font=("Helvetica", 12))
    
    btn_limpar_filtro.pack(side="left", padx=10, pady=10)

    btn_limpar_filtro.bind(
    "<Enter>",
    lambda e: btn_limpar_filtro.config(fg="white", bg="#8B8989"))  



    btn_limpar_filtro.bind(
    "<Leave>",
    lambda e: btn_limpar_filtro.config(fg="black", bg="#CDC9C9")) 


    # ------------------------- FRAME LISTA -----------------------------

    frame_lista = tk.Frame(janela_principal, pady=20)

    frame_lista.pack(fill="both", expand=True)

    global tree_produtos


    tree_produtos = ttk.Treeview(frame_lista, 
                                 columns=("Codigo", "Nome", "Preço", "Quantidade", "Localizacao", "ID"),
                                 show="headings")
    



    scroll_y = ttk.Scrollbar(frame_lista, orient="vertical", command=tree_produtos.yview)



    tree_produtos.configure(yscrollcommand=scroll_y.set)


    scroll_y.pack(side="right", fill="y")



    tree_produtos.heading("Codigo", text="Código")

    tree_produtos.heading("Nome", text="Nome")

    tree_produtos.heading("Preço", text="Preço")

    tree_produtos.heading("Quantidade", text="Quantidade")

    tree_produtos.heading("Localizacao", text="Localização")

    tree_produtos.heading("ID", text="ID")





    tree_produtos.column("Codigo", width=200, anchor="center")

    tree_produtos.column("Nome", width=200, anchor="center")

    tree_produtos.column("Preço", width=100, anchor="center")

    tree_produtos.column("Quantidade", width=100, anchor="center")

    tree_produtos.column("Localizacao", width=150, anchor="center")

    tree_produtos.column("ID", width=100, anchor="center")


    tree_produtos.pack(fill="both", expand=True)


    tree_produtos.bind('<<TreeviewSelect>>', preencher_campos_produto)


    carregar_dados()



    janela_principal.mainloop()


    



                                 
                            




















# ====== INTERFACE DA PÁGINA DE LOGIN ======

janela_login = tk.Tk()

janela_login.title("Login - Controle de estoque")

janela_login.state('zoomed')

janela_login.resizable(False, False)



def caminho_absoluto(relativo):

    if hasattr(sys, '_MEIPASS'):

        return os.path.join(sys._MEIPASS, relativo)
    
    return os.path.join(os.path.abspath("."), relativo)


icone = PhotoImage(file=caminho_absoluto("logo.png"))

janela_login.iconphoto(True, icone)





# ------ IMAGEM DE FUNDO ------

imagem_original = Image.open(caminho_absoluto("background.png"))


label_fundo = tk.Label(janela_login)

label_fundo.place(x=0, y=0, relwidth=1, relheight=1)



def atualizar_fundo(event=None):

    largura_janela = janela_login.winfo_width()

    altura_janela = janela_login.winfo_height()


    if largura_janela < 10 or altura_janela < 10:
        return
    

    largura_img, altura_img = imagem_original.size


    escala = max(largura_janela / largura_img, altura_janela / altura_img)

    nova_largura = int(largura_img * escala)

    nova_altura = int(altura_img * escala)



    imagem_redimensionada = imagem_original.resize(
        (nova_largura, nova_altura),
        Image.LANCZOS

    )



    esquerda = (nova_largura - largura_janela) // 2
    topo = (nova_altura - altura_janela) // 2
    direita = esquerda + largura_janela
    baixo = topo + altura_janela



    imagem_cortada = imagem_redimensionada.crop((esquerda, topo, direita, baixo))

    overlay = Image.new("RGBA", imagem_cortada.size, (0, 0, 0, 90))



    imagem_final = Image.alpha_composite(
        imagem_cortada.convert("RGBA"),
        overlay

    )

    imagem_tk = ImageTk.PhotoImage(imagem_final)

    label_fundo.config(image=imagem_tk)

    label_fundo.imagem = imagem_tk


janela_login.after(100, atualizar_fundo)

janela_login.bind("<Configure>", atualizar_fundo)



# ------ FRAME CENTRAL DA PÁGINA DE LOGIN ------

frame_login = tk.Frame(janela_login, padx=40, pady=40, bg="#F8F9FA")

frame_login.place(relx=0.5, rely=0.5, anchor='center')


lbl_titulo_login = tk.Label(frame_login, text="Controle de estoque", bg="#F8F9FA", font=("Helvetica", 20, "bold"))
lbl_titulo_login.grid(row=0, column=0, columnspan=2, pady=20)




tk.Label(frame_login, text="Usuário", bg="#F8F9FA", font=("Helvetica", 14)).grid(row=1, column=0, pady=10, sticky='e')

entrada_usuario = tk.Entry(frame_login, font=("Helvetica", 14))
entrada_usuario.grid(row=1, column=1, pady=10)




tk.Label(frame_login, text="Senha", bg="#F8F9FA", font=("Helvetica", 14)).grid(row=2, column=0, pady=10, sticky='e')

entrada_senha = tk.Entry(frame_login, font=("Helvetica", 14))
entrada_senha.grid(row=2, column=1, pady=10)











btn_login = tk.Button(frame_login,
                      text="Login",
                      command=verificar_login,
                      bg="#2A9D8F",
                      fg="white",
                      font=("Helvetica", 14),
                      width=15)

btn_login.grid(row=3, column=0, columnspan=2, pady=10)









btn_cadastrar_usuario = tk.Button(frame_login,
                      text="Cadastrar Usuário",
                      command=abrir_tela_usuarios,
                      bg="#264653",
                      fg="white",
                      font=("Helvetica", 14),
                      width=15)

btn_cadastrar_usuario.grid(row=4, column=0, columnspan=2, pady=10)



criar_usuario_admin()









janela_login.mainloop()
